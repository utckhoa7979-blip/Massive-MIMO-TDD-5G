import numpy as np
import scipy
import matplotlib.pyplot as plt
import pandas as pd
from pathlib import Path


# ==========================================================
# 1. THAM SO CHUNG CUA HE THONG MASSIVE MIMO TDD
# ==========================================================

K = 4
M_values = np.arange(10, 101, 10)

tau_c = 200

# Cong suat phat UE: 100 mW = 0.1 W
P_UE_W = 0.1

# Tham so vo tuyen
fc_GHz = 3.5
bandwidth_Hz = 20e6
noise_figure_dB = 7.0

# Tham so mo hinh kenh
cell_radius_m = 300
min_distance_m = 50
h_BS_m = 25
h_UE_m = 1.5
shadow_fading_std_dB = 6.0

# Cau hinh Monte Carlo
num_UE_configurations = 20
num_channel_realizations = 20


# ==========================================================
# 2. BA KICH BAN MO PHONG
# ==========================================================

scenarios = {
    "KB1_SingleCell": {
        "L": 1,
        "K": K,
        "f": 1,
        "tau_p": 4,
    },
    "KB2_MultiCell_f1": {
        "L": 4,
        "K": K,
        "f": 1,
        "tau_p": 4,
    },
    "KB3_MultiCell_f4": {
        "L": 4,
        "K": K,
        "f": 4,
        "tau_p": 16,
    },
}


# ==========================================================
# 3. KHOI TAO VI TRI BS VA UE
# ==========================================================

def generate_bs_positions(L, cell_radius_m):
    """
    Tao vi tri BS.
    L=1: mot cell.
    L=4: bo tri BS theo mang luoi 2x2.
    """

    if L == 1:
        return np.array([[0.0, 0.0]])

    if L == 4:
        spacing = 2.0 * cell_radius_m

        return np.array([
            [0.0, 0.0],
            [spacing, 0.0],
            [0.0, spacing],
            [spacing, spacing],
        ])

    raise ValueError("Chuong trinh hien ho tro L=1 hoac L=4.")


def generate_ue_positions(
    bs_positions, K, cell_radius_m, min_distance_m, rng
):
    """
    Tao K UE trong moi cell.
    UE phan bo deu theo dien tich hinh tron.
    Khoang cach UE den BS phuc vu nam trong [dmin, Rcell].
    """

    L = len(bs_positions)

    if not 0 <= min_distance_m < cell_radius_m:
        raise ValueError("Khoang cach UE toi thieu khong hop le.")

    # Goc cua UE
    angles = rng.uniform(
        0.0, 2.0 * np.pi, size=(L, K)
    )

    # Lay mau ban kinh de UE phan bo deu theo dien tich
    radii = np.sqrt(
        rng.uniform(
            min_distance_m**2,
            cell_radius_m**2,
            size=(L, K),
        )
    )

    # Toa do UE, kich thuoc (L, K, 2)
    offsets = np.stack(
        [
            radii * np.cos(angles),
            radii * np.sin(angles),
        ],
        axis=-1,
    )

    return bs_positions[:, None, :] + offsets


# ==========================================================
# 4. MO HINH SUY HAO UMa-NLOS VA HE SO BETA
# ==========================================================

def compute_large_scale_coefficients(
    bs_positions, ue_positions, rng
):
    """
    Tinh suy hao UMa-NLOS, shadow fading va beta.

    beta[j, l, k]:
        He so kenh quy mo lon tu UE k o cell l den BS j.

    Kich thuoc beta: (L, L, K)
    """

    L = len(bs_positions)
    K_local = ue_positions.shape[1]

    # Toc do anh sang (m/s)
    c_mps = 299_792_458.0

    # Gia thiet chieu cao moi truong hE
    h_E_m = 1.0

    # Chieu cao anten hieu dung
    h_BS_eff = h_BS_m - h_E_m
    h_UE_eff = h_UE_m - h_E_m

    if h_BS_eff <= 0 or h_UE_eff <= 0:
        raise ValueError("Chieu cao anten hieu dung khong hop le.")

    # Khoang cach diem gay, fc tinh bang Hz
    d_BP_m = (
        4.0 * h_BS_eff * h_UE_eff
        * (fc_GHz * 1e9) / c_mps
    )

    beta = np.zeros((L, L, K_local))
    d2D_all = np.zeros((L, L, K_local))
    pathloss_total = np.zeros((L, L, K_local))

    for j in range(L):          # BS dang thu
        for l in range(L):      # Cell chua UE

            dx = (
                ue_positions[l, :, 0]
                - bs_positions[j, 0]
            )
            dy = (
                ue_positions[l, :, 1]
                - bs_positions[j, 1]
            )

            # Khoang cach 2D va 3D
            d2D = np.sqrt(dx**2 + dy**2)

            d3D = np.sqrt(
                d2D**2 + (h_BS_m - h_UE_m)**2
            )

            d2D_all[j, l, :] = d2D

            # UMa-LOS tham chieu
            PL1 = (
                28.0
                + 22.0 * np.log10(d3D)
                + 20.0 * np.log10(fc_GHz)
            )

            PL2 = (
                28.0
                + 40.0 * np.log10(d3D)
                + 20.0 * np.log10(fc_GHz)
                - 9.0 * np.log10(
                    d_BP_m**2
                    + (h_BS_m - h_UE_m)**2
                )
            )

            PL_LOS = np.where(
                d2D < d_BP_m,
                PL1,
                PL2,
            )

            # UMa-NLOS tham chieu
            PL_NLOS_prime = (
                13.54
                + 39.08 * np.log10(d3D)
                + 20.0 * np.log10(fc_GHz)
                - 0.6 * (h_UE_m - 1.5)
            )

            # Suy hao NLOS cuoi cung
            PL_NLOS = np.maximum(
                PL_LOS,
                PL_NLOS_prime,
            )

            # Shadow fading Gaussian, sigma = 6 dB
            shadow_fading = rng.normal(
                loc=0.0,
                scale=shadow_fading_std_dB,
                size=K_local,
            )

            PL_total = PL_NLOS + shadow_fading
            pathloss_total[j, l, :] = PL_total

            # Chuyen suy hao dB sang he so tuyen tinh
            beta[j, l, :] = 10.0**(-PL_total / 10.0)

    return beta, d2D_all, pathloss_total


# ==========================================================
# 5. SINH KENH TRUYEN RAYLEIGH
# ==========================================================

def generate_rayleigh_channels(beta, M, rng):
    """
    Sinh kenh Rayleigh phuc, khong tuong quan theo khong gian.

    beta[j, l, k]: he so kenh tu UE k o cell l den BS j.
    M: so anten tai moi BS.

    Dau ra:
        H[j, l, k, m] la kenh tu UE k o cell l
        den anten m cua BS j.

    Kich thuoc H: (L, L, K, M)
    """

    if M <= 0:
        raise ValueError("So anten M phai lon hon 0.")

    if not np.all(np.isfinite(beta)) or np.any(beta <= 0):
        raise ValueError("Beta phai la cac gia tri huu han va duong.")

    shape = beta.shape + (M,)

    # Gaussian phuc CN(0,1)
    z = (
        rng.standard_normal(size=shape)
        + 1j * rng.standard_normal(size=shape)
    ) / np.sqrt(2.0)

    # Dieu chinh phuong sai theo beta
    H = np.sqrt(beta[..., np.newaxis]) * z

    return H
# ==========================================================
# 6. TAO CHUOI PILOT
# ==========================================================

def generate_pilot_sequences(L, K, f, tau_p):
    """
    Tao chuoi pilot truc giao theo he so tai su dung f.

    Dau ra co kich thuoc (L, K, tau_p).
    Moi chuoi pilot co binh phuong chuan bang tau_p.
    """
    if f < 1:
        raise ValueError("He so tai su dung pilot f phai >= 1.")

    if tau_p != f * K:
        raise ValueError("Can dam bao tau_p = f * K.")

    n = np.arange(tau_p)
    pilots = np.empty((L, K, tau_p), dtype=complex)

    for l in range(L):
        for k in range(K):
            # Cac cell cung nhom dung lai tap pilot
            pilot_id = (l % f) * K + k

            pilots[l, k, :] = np.exp(
                1j * 2.0 * np.pi * pilot_id * n / tau_p
            )

    return pilots
# ==========================================================
# 7. MO PHONG TIN HIEU PILOT THU TAI BS
# ==========================================================

def simulate_received_pilots(
    H, pilots, transmit_power_W, noise_power_W, rng
):
    """
    H: (L, L, K, M)
    pilots: (L, K, tau_p)
    Tra ve Y_p co kich thuoc (L, M, tau_p).
    """
    L, _, K_local, M = H.shape
    tau_p = pilots.shape[2]

    Y_p = np.zeros((L, M, tau_p), dtype=complex)

    for j in range(L):
        for l in range(L):
            for k in range(K_local):
                Y_p[j] += (
                    np.sqrt(transmit_power_W)
                    * np.outer(H[j, l, k, :], pilots[l, k, :])
                )

        noise = np.sqrt(noise_power_W / 2.0) * (
            rng.standard_normal((M, tau_p))
            + 1j * rng.standard_normal((M, tau_p))
        )

        Y_p[j] += noise

    return Y_p


def correlate_pilot_observations(Y_p, pilots):
    """
    Tuong quan tin hieu pilot thu voi tung chuoi pilot.
    Tra ve mang y_obs co kich thuoc (L, L, K, M).
    """
    L, M, _ = Y_p.shape
    K_local = pilots.shape[1]

    y_obs = np.zeros(
        (L, L, K_local, M),
        dtype=complex
    )

    for j in range(L):
        for l in range(L):
            for k in range(K_local):
                y_obs[j, l, k, :] = (
                    Y_p[j] @ np.conj(pilots[l, k, :])
                )

    return y_obs
# ==========================================================
# 8. UOC LUONG KENH MMSE
# ==========================================================

def estimate_channels_mmse(
    beta, pilots, y_obs, pilot_power_W, noise_power_W
):
    """
    Uoc luong kenh MMSE cho kenh Rayleigh khong tuong quan.

    beta:   (L, L, K)
    pilots: (L, K, tau_p)
    y_obs:  (L, L, K, M)

    Ket qua H_hat: (L, L, K, M)
    """
    L, _, K_local = beta.shape
    M = y_obs.shape[-1]
    tau_p = pilots.shape[-1]

    H_hat = np.zeros(
        (L, L, K_local, M),
        dtype=complex
    )

    for j in range(L):
        for l in range(L):
            for k in range(K_local):

                # Phuong sai tong hop cua quan sat pilot
                psi = noise_power_W

                # Cong cac UE su dung cung chuoi pilot
                for l2 in range(L):
                    for k2 in range(K_local):
                        if np.allclose(
                            pilots[l2, k2, :],
                            pilots[l, k, :]
                        ):
                            psi += (
                                tau_p
                                * pilot_power_W
                                * beta[j, l2, k2]
                            )

                # Uoc luong MMSE cho UE k tai cell l, BS j
                H_hat[j, l, k, :] = (
                    np.sqrt(pilot_power_W)
                    * beta[j, l, k]
                    / psi
                    * y_obs[j, l, k, :]
                )

    return H_hat
# ==========================================================
# 9. XAY DUNG CAC BO KET HOP MR, ZF VA M-MMSE
# ==========================================================

def compute_combining_vectors(
    H_hat, beta, pilots, method,
    bs_index, transmit_power_W, noise_power_W
):
    """
    Tao vector ket hop cho cac UE trong cell phuc vu.

    H_hat: (L, L, K, M)
    beta:  (L, L, K)
    pilots: (L, K, tau_p)

    Dau ra V co kich thuoc (M, K).
    Moi cot la vector ket hop cua mot UE trong cell.
    """

    L, _, K_local, M = H_hat.shape
    tau_p = pilots.shape[-1]

    # Ma tran kenh uoc luong cua cac UE trong cell phuc vu
    H_serving = H_hat[bs_index, bs_index, :, :].T

    # ------------------------------------------------------
    # MR: dung truc tiep kenh uoc luong
    # ------------------------------------------------------
    if method == "MR":
        return H_serving.copy()

    # ------------------------------------------------------
    # ZF: giam can nhieu noi cell
    # ------------------------------------------------------
    if method == "ZF":
        gram_matrix = (
            H_serving.conj().T @ H_serving
        )

        return H_serving @ np.linalg.pinv(gram_matrix)

    # ------------------------------------------------------
    # M-MMSE: su dung thong tin kenh cua toan mang
    # ------------------------------------------------------
    if method == "M-MMSE":
        identity = np.eye(M, dtype=complex)

        A = noise_power_W * identity

        for l in range(L):
            for k in range(K_local):

                # Tinh psi cho nhom UE dung chung pilot
                psi = noise_power_W

                for l2 in range(L):
                    for k2 in range(K_local):
                        if np.allclose(
                            pilots[l2, k2, :],
                            pilots[l, k, :]
                        ):
                            psi += (
                                tau_p
                                * transmit_power_W
                                * beta[bs_index, l2, k2]
                            )

                beta_lk = beta[bs_index, l, k]

                # Phuong sai sai so uoc luong MMSE
                C_error = (
                    beta_lk
                    - tau_p
                    * transmit_power_W
                    * beta_lk**2
                    / psi
                )

                C_error = max(float(C_error), 0.0)

                h_est = H_hat[bs_index, l, k, :]

                A += transmit_power_W * (
                    np.outer(h_est, h_est.conj())
                    + C_error * identity
                )

        # Moi cot tuong ung mot UE trong cell phuc vu
        V = np.linalg.solve(
            A,
            transmit_power_W * H_serving
        )

        return V

    raise ValueError(
        "Phuong phap phai la MR, ZF hoac M-MMSE."
    )
# ==========================================================
# 10. CAC DAI LUONG PHUC VU TINH SINR UATF
# ==========================================================

def compute_uplink_uatf_sample_terms(
    H, V, bs_index, transmit_power_W, noise_power_W
):
    """
    Tinh cac dai luong cua mot mau kenh de tong hop SINR UatF.

    H: (L, L, K, M) - kenh thuc
    V: (M, K) - cac vector ket hop tai BS bs_index

    Tra ve:
        desired_coeff: he so kenh mong muon v^H h
        total_rx_power: tong cong suat tin hieu sau ket hop
        noise_term: cong suat nhieu sau ket hop

    Cac dai luong nay se duoc lay trung binh qua Monte Carlo.
    """

    L, _, K_local, M = H.shape

    desired_coeff = np.zeros(K_local, dtype=complex)
    total_rx_power = np.zeros(K_local, dtype=float)
    noise_term = np.zeros(K_local, dtype=float)

    for k in range(K_local):
        v = V[:, k]

        # He so kenh mong muon sau ket hop
        desired_coeff[k] = np.vdot(
            v, H[bs_index, bs_index, k, :]
        )

        # Tong cong suat cua tat ca UE sau ket hop
        for l in range(L):
            for i in range(K_local):
                effective_channel = np.vdot(
                    v, H[bs_index, l, i, :]
                )

                total_rx_power[k] += (
                    transmit_power_W
                    * np.abs(effective_channel) ** 2
                )

        # Cong suat nhieu sau ket hop
        noise_term[k] = (
            noise_power_W * np.vdot(v, v).real
        )

    return desired_coeff, total_rx_power, noise_term
# ==========================================================
# 11. TINH SINR UatF VA HIEU SUAT PHO (SE)
# ==========================================================

def compute_uatf_sinr_and_se(
    mean_desired_coeff,
    mean_total_rx_power,
    mean_noise_power,
    transmit_power_W,
    tau_c,
    tau_p
):
    """
    Tinh SINR theo Use-and-then-forget (UatF)
    va spectral efficiency (SE).

    Cac dau vao la gia tri trung binh qua Monte Carlo.
    """

    # Cong suat tin hieu mong muon
    signal_power = (
        transmit_power_W
        * np.abs(mean_desired_coeff) ** 2
    )

    # Tong cong suat nhieu va nhieu dong kenh
    interference_plus_noise = (
        mean_total_rx_power
        - signal_power
        + mean_noise_power
    )

    # Dam bao mau so khong am do sai so so hoc
    interference_plus_noise = np.maximum(
        interference_plus_noise,
        np.finfo(float).tiny
    )

    # SINR theo UatF
    sinr = signal_power / interference_plus_noise

    # He so tai pilot va du lieu trong khoi tuong hop
    prelog_factor = (tau_c - tau_p) / tau_c

    # Hieu suat pho, don vi bit/s/Hz
    se = prelog_factor * np.log2(1 + sinr)

    return sinr, se
# ==========================================================
# 12. MO PHONG MONTE CARLO DAY DU
# ==========================================================

def run_monte_carlo_simulation(noise_power_W):
    """Monte Carlo; luu ket qua rieng cho tung cau hinh UE."""

    rng_mc = np.random.default_rng(2026)
    methods = ("MR", "ZF", "M-MMSE")
    all_detail_records = []
    summary_results = {}
    max_M = int(np.max(M_values))

    for name, cfg in scenarios.items():
        print("\n" + "=" * 60)
        print(f"MONTE CARLO: {name}")
        print("=" * 60)

        L = cfg["L"]
        K_local = cfg["K"]
        tau_p = cfg["tau_p"]
        bs_positions = generate_bs_positions(L, cell_radius_m)
        pilots = generate_pilot_sequences(L, K_local, cfg["f"], tau_p)
        scenario_records = []

        for config_idx in range(num_UE_configurations):
            ue_positions = generate_ue_positions(
                bs_positions, K_local, cell_radius_m,
                min_distance_m, rng_mc
            )

            # Beta va shadow fading duoc giu co dinh trong cau hinh UE nay.
            beta, _, _ = compute_large_scale_coefficients(
                bs_positions, ue_positions, rng_mc
            )

            # Tich luy cac mau fading cho cau hinh UE hien tai.
            config_acc = {
                int(M): {
                    method: {
                        "desired": np.zeros((L, K_local), dtype=complex),
                        "total_rx": np.zeros((L, K_local), dtype=float),
                        "noise": np.zeros((L, K_local), dtype=float),
                    }
                    for method in methods
                }
                for M in M_values
            }

            for realization_idx in range(num_channel_realizations):
                H_max = generate_rayleigh_channels(beta, max_M, rng_mc)
                Y_p_max = simulate_received_pilots(
                    H_max, pilots, P_UE_W, noise_power_W, rng_mc
                )

                for M_value in M_values:
                    M_int = int(M_value)
                    H = H_max[..., :M_int]
                    Y_p = Y_p_max[:, :M_int, :]
                    y_obs = correlate_pilot_observations(Y_p, pilots)
                    H_hat = estimate_channels_mmse(
                        beta, pilots, y_obs, P_UE_W, noise_power_W
                    )

                    for bs_index in range(L):
                        for method in methods:
                            V = compute_combining_vectors(
                                H_hat, beta, pilots, method,
                                bs_index, P_UE_W, noise_power_W
                            )
                            desired, total_rx, noise = (
                                compute_uplink_uatf_sample_terms(
                                    H, V, bs_index, P_UE_W, noise_power_W
                                )
                            )
                            acc = config_acc[M_int][method]
                            acc["desired"][bs_index, :] += desired
                            acc["total_rx"][bs_index, :] += total_rx
                            acc["noise"][bs_index, :] += noise

            # Tinh SINR UatF trong cau hinh UE nay sau khi lay trung binh fading.
            for M_value in M_values:
                M_int = int(M_value)
                record = {
                    "Scenario": name,
                    "Configuration": config_idx + 1,
                    "M": M_int,
                }

                for method in methods:
                    acc = config_acc[M_int][method]
                    mean_desired = acc["desired"] / num_channel_realizations
                    mean_total_rx = acc["total_rx"] / num_channel_realizations
                    mean_noise = acc["noise"] / num_channel_realizations

                    sinr, se = compute_uatf_sinr_and_se(
                        mean_desired, mean_total_rx, mean_noise,
                        P_UE_W, tau_c, tau_p
                    )

                    if not (np.isfinite(sinr).all() and np.isfinite(se).all()):
                        raise FloatingPointError(
                            f"SINR/SE khong huu han: {name}, M={M_int}, {method}"
                        )

                    # Trung binh tren cac UE cua cac BS; day la mot diem du lieu cua cau hinh UE.
                    record[f"{method}_SINR"] = float(np.mean(sinr))
                    record[f"{method}_SE"] = float(np.mean(se))

                scenario_records.append(record)
                all_detail_records.append(record.copy())

            print(
                f"  UE configuration {config_idx + 1}/"
                f"{num_UE_configurations} completed"
            )

        # Thong ke tren 20 cau hinh vi tri UE, rieng cho tung M va phuong phap.
        summary_results[name] = {}
        for M_value in M_values:
            M_int = int(M_value)
            rows_m = [r for r in scenario_records if r["M"] == M_int]
            summary_results[name][M_int] = {}
            for method in methods:
                se_values = np.array([r[f"{method}_SE"] for r in rows_m], dtype=float)
                sinr_values = np.array([r[f"{method}_SINR"] for r in rows_m], dtype=float)
                summary_results[name][M_int][method] = {
                    "mean_SE": float(np.mean(se_values)),
                    "std_SE": float(np.std(se_values, ddof=1)) if len(se_values) > 1 else 0.0,
                    "min_SE": float(np.min(se_values)),
                    "max_SE": float(np.max(se_values)),
                    "mean_SINR": float(np.mean(sinr_values)),
                    "std_SINR": float(np.std(sinr_values, ddof=1)) if len(sinr_values) > 1 else 0.0,
                    "min_SINR": float(np.min(sinr_values)),
                    "max_SINR": float(np.max(sinr_values)),
                }

                s = summary_results[name][M_int][method]
                print(
                    f"M={M_int:3d} | {method:6s} | "
                    f"Mean SE={s['mean_SE']:.4f} +/- {s['std_SE']:.4f} bit/s/Hz | "
                    f"Mean SINR={s['mean_SINR']:.4f} +/- {s['std_SINR']:.4f}"
                )

    return {"summary": summary_results, "detail_records": all_detail_records}


def export_results_to_excel(mc_data, output_path=None):
    """Xuat cac bang chi tiet, thong ke va so sanh vao mot file Excel co dinh dang."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if output_path is None:
        output_path = Path(__file__).resolve().parent / "MIMOTDD_MonteCarlo_Results.xlsx"
    else:
        output_path = Path(output_path)

    detail_df = pd.DataFrame(mc_data["detail_records"])
    summary_data = mc_data["summary"]
    methods = ("MR", "ZF", "M-MMSE")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Moi kich ban co mot bang 20 cau hinh UE x 30 cot SE, va mot bang SINR.
        for scenario_name in scenarios:
            scen = detail_df[detail_df["Scenario"] == scenario_name].copy()
            scenario_short = scenario_name.split("_")[0]

            for metric in ("SE", "SINR"):
                wide = {"Cau_hinh_UE": list(range(1, num_UE_configurations + 1))}
                for M_value in M_values:
                    M_int = int(M_value)
                    for method in methods:
                        col = f"{method}_{metric}"
                        subset = scen[scen["M"] == M_int].set_index("Configuration")
                        lookup = subset[col].to_dict()
                        wide[f"M{M_int}_{method}"] = [lookup.get(c, np.nan) for c in wide["Cau_hinh_UE"]]
                pd.DataFrame(wide).to_excel(
                    writer,
                    sheet_name=f"{scenario_short}_{metric}_20cauHinh",
                    index=False,
                )

        # Bang thong ke: mean/std/min/max giua 20 cau hinh UE.
        stats_rows = []
        for scenario_name in scenarios:
            for M_value in M_values:
                M_int = int(M_value)
                for method in methods:
                    stat = summary_data[scenario_name][M_int][method]
                    stats_rows.append({
                        "Kich_ban": scenario_name,
                        "M": M_int,
                        "Phuong_phap": method,
                        "SE_trung_binh": stat["mean_SE"],
                        "SE_do_lech_chuan": stat["std_SE"],
                        "SE_nho_nhat": stat["min_SE"],
                        "SE_lon_nhat": stat["max_SE"],
                        "SINR_trung_binh": stat["mean_SINR"],
                        "SINR_do_lech_chuan": stat["std_SINR"],
                        "SINR_nho_nhat": stat["min_SINR"],
                        "SINR_lon_nhat": stat["max_SINR"],
                    })
        pd.DataFrame(stats_rows).to_excel(writer, sheet_name="Thong_ke_20_cau_hinh", index=False)

        # Bang so sanh SE: moi hang la M; cot la kich ban + phuong phap.
        compare_rows = []
        for M_value in M_values:
            row = {"M": int(M_value)}
            for scenario_name in scenarios:
                for method in methods:
                    stat = summary_data[scenario_name][int(M_value)][method]
                    row[f"{scenario_name}_{method}_MeanSE"] = stat["mean_SE"]
                    row[f"{scenario_name}_{method}_StdSE"] = stat["std_SE"]
            compare_rows.append(row)
        pd.DataFrame(compare_rows).to_excel(writer, sheet_name="So_sanh_tong_the", index=False)

        # Bang tong quan o M=100, de tham chieu nhanh trong phan ket luan.
        rows_100 = []
        M_focus = int(np.max(M_values))
        for scenario_name in scenarios:
            for method in methods:
                stat = summary_data[scenario_name][M_focus][method]
                rows_100.append({
                    "Kich_ban": scenario_name,
                    "M": M_focus,
                    "Phuong_phap": method,
                    "SE_trung_binh": stat["mean_SE"],
                    "SE_do_lech_chuan": stat["std_SE"],
                    "SINR_trung_binh": stat["mean_SINR"],
                    "SINR_do_lech_chuan": stat["std_SINR"],
                })
        pd.DataFrame(rows_100).to_excel(writer, sheet_name="So_sanh_M100", index=False)

        # Du lieu dang dai de truy vet / xu ly lai sau nay.
        detail_df.sort_values(["Scenario", "Configuration", "M"]).to_excel(
            writer, sheet_name="Du_lieu_chi_tiet", index=False
        )

        # Thong so mo phong va so mau.
        config_rows = [
            {"Thong_so": "So UE moi cell (K)", "Gia_tri": K},
            {"Thong_so": "So anten BS (M)", "Gia_tri": ", ".join(map(str, M_values.tolist()))},
            {"Thong_so": "Khoi tuong hop tau_c", "Gia_tri": tau_c},
            {"Thong_so": "Cong suat phat UE (W)", "Gia_tri": P_UE_W},
            {"Thong_so": "Tan so song mang (GHz)", "Gia_tri": fc_GHz},
            {"Thong_so": "Bang thong (MHz)", "Gia_tri": bandwidth_Hz / 1e6},
            {"Thong_so": "Noise figure (dB)", "Gia_tri": noise_figure_dB},
            {"Thong_so": "Ban kinh cell (m)", "Gia_tri": cell_radius_m},
            {"Thong_so": "Khoang cach UE toi thieu (m)", "Gia_tri": min_distance_m},
            {"Thong_so": "Chieu cao BS (m)", "Gia_tri": h_BS_m},
            {"Thong_so": "Chieu cao UE (m)", "Gia_tri": h_UE_m},
            {"Thong_so": "Do lech chuan shadow fading (dB)", "Gia_tri": shadow_fading_std_dB},
            {"Thong_so": "So cau hinh vi tri UE", "Gia_tri": num_UE_configurations},
            {"Thong_so": "So mau kenh moi cau hinh", "Gia_tri": num_channel_realizations},
            {"Thong_so": "Tong mau kenh moi kich ban va M", "Gia_tri": num_UE_configurations * num_channel_realizations},
            {"Thong_so": "Phuong phap ket hop", "Gia_tri": "MR, ZF, M-MMSE"},
        ]
        pd.DataFrame(config_rows).to_excel(writer, sheet_name="Thong_so_mo_phong", index=False)

        # Dinh dang workbook: tieu de, loc, co dinh hang dau va do rong cot.
        wb = writer.book
        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="203040")
        thin_gray = Side(style="thin", color="D9E2F3")

        for ws in wb.worksheets:
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            ws.row_dimensions[1].height = 30
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(bottom=thin_gray)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
                    cell.alignment = Alignment(vertical="center")
                    if isinstance(cell.value, (float, np.floating)):
                        cell.number_format = "0.0000"
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                values = [str(c.value) if c.value is not None else "" for c in list(col_cells)[:100]]
                max_len = max((len(v) for v in values), default=10)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 27)

    print(f"\nDa xuat file Excel: {output_path}")
    print("Cac sheet chinh: 20 cau hinh UE theo tung kich ban, thong ke mean/std, so sanh tong the, M=100, va du lieu day du.")
    return output_path

# ==========================================================
# 6. HAM MAIN: KIEM TRA MO HINH VA KENH
# ==========================================================

def main():
    print("MASSIVE MIMO TDD - UPLINK SIMULATION")
    print("-" * 50)

    print(f"NumPy version: {np.__version__}")
    print(f"SciPy version: {scipy.__version__}")
    print(f"M values: {M_values.tolist()}")
    print(f"Coherence block: {tau_c}")
    print(f"UE transmit power: {P_UE_W} W")
    print(f"Bandwidth: {bandwidth_Hz / 1e6:.1f} MHz")
    print(f"Noise figure: {noise_figure_dB} dB")

    # Dung seed de ket qua kiem tra co the lap lai
    rng = np.random.default_rng(42)
        # Cong suat nhieu thu
    noise_power_dBm = (
        -174
        + 10 * np.log10(bandwidth_Hz)
        + noise_figure_dB
    )

    noise_power_W = 10 ** (
        (noise_power_dBm - 30) / 10
    )

    print(f"Noise power: {noise_power_dBm:.2f} dBm")
    print(f"Noise power: {noise_power_W:.3e} W")

    print("\nSIMULATION SCENARIOS")

    for name, cfg in scenarios.items():
        print(
            f"{name}: L={cfg['L']}, "
            f"K={cfg['K']}, f={cfg['f']}, "
            f"tau_p={cfg['tau_p']}"
        )

    print("\nNETWORK, PATHLOSS AND RAYLEIGH CHANNEL TEST")
    print("-" * 50)

    # Duyet qua ba kich ban
    for name, cfg in scenarios.items():

        # 1. Khoi tao vi tri BS va UE
        bs_positions = generate_bs_positions(
            cfg["L"],
            cell_radius_m,
        )

        ue_positions = generate_ue_positions(
            bs_positions,
            cfg["K"],
            cell_radius_m,
            min_distance_m,
            rng,
        )

        # 2. Tinh khoang cach UE den BS phuc vu
        serving_distances = np.linalg.norm(
            ue_positions - bs_positions[:, None, :],
            axis=2,
        )

        # 3. Tinh beta va suy hao duong truyen
        beta, d2D_all, pathloss_total = (
            compute_large_scale_coefficients(
                bs_positions,
                ue_positions,
                rng,
            )
        )

        # 4. Sinh kenh Rayleigh voi M = 10 anten de kiem tra
        H = generate_rayleigh_channels(
            beta,
            M=10,
            rng=rng,
        )

        # 5. Cong suat kenh chuan hoa theo beta
        normalized_power = np.mean(
            np.abs(H)**2 / beta[..., np.newaxis]
        )

        # 6. In ket qua kiem tra
        print(f"\n{name}")
        print(f"BS position array shape: {bs_positions.shape}")
        print(f"UE position array shape: {ue_positions.shape}")

        print(
            "Serving distance range: "
            f"{serving_distances.min():.1f} - "
            f"{serving_distances.max():.1f} m"
        )

        print(f"Beta array shape: {beta.shape}")

        print(
            "Pathloss range: "
            f"{pathloss_total.min():.1f} - "
            f"{pathloss_total.max():.1f} dB"
        )

        print(
            "Beta range: "
            f"{beta.min():.3e} - {beta.max():.3e}"
        )

        print(f"Rayleigh channel shape: {H.shape}")
        print(f"Complex-valued: {np.iscomplexobj(H)}")
        print(f"All finite: {np.isfinite(H).all()}")

        print(
            "Mean normalized channel power "
            f"(expected approximately 1): {normalized_power:.3f}"
        )
        # 7. TAO VA KIEM TRA CHUOI PILOT
        pilots = generate_pilot_sequences(
            L=cfg["L"],
            K=cfg["K"],
            f=cfg["f"],
            tau_p=cfg["tau_p"],
        )

        flat_pilots = pilots.reshape(
            -1, cfg["tau_p"]
        )

        unique_pilots = np.unique(
            np.round(flat_pilots, 8),
            axis=0,
        )

        gram = (
            unique_pilots
            @ unique_pilots.conj().T
        )

        off_diagonal = (
            gram - np.diag(np.diag(gram))
        )

        max_cross_correlation = (
            np.max(np.abs(off_diagonal))
            if len(unique_pilots) > 1
            else 0.0
        )

        print(f"Pilot array shape: {pilots.shape}")
        print(
            f"Number of unique pilots: "
            f"{len(unique_pilots)}"
        )
        print(
            "Maximum cross-correlation: "
            f"{max_cross_correlation:.3e}"
        )
                # Mo phong tin hieu pilot thu tai BS
        Y_p = simulate_received_pilots(
            H,
            pilots,
            P_UE_W,
            noise_power_W,
            rng
        )

        # Tuong quan tin hieu pilot
        y_obs = correlate_pilot_observations(
            Y_p,
            pilots
        )
                # Uoc luong kenh MMSE
        H_hat = estimate_channels_mmse(
            beta,
            pilots,
            y_obs,
            P_UE_W,
            noise_power_W
        )
                # Kiem tra ba bo ket hop tai BS dau tien
        V_MR = compute_combining_vectors(
            H_hat, beta, pilots, "MR",
            0, P_UE_W, noise_power_W
        )

        V_ZF = compute_combining_vectors(
            H_hat, beta, pilots, "ZF",
            0, P_UE_W, noise_power_W
        )

        V_MMMSE = compute_combining_vectors(
            H_hat, beta, pilots, "M-MMSE",
            0, P_UE_W, noise_power_W
        )
                # Kiem tra cac dai luong phuc vu tinh SINR
        desired_mr, total_mr, noise_mr = (
            compute_uplink_uatf_sample_terms(
                H, V_MR, 0, P_UE_W, noise_power_W
            )
        )

        print(f"Desired coefficient shape: {desired_mr.shape}")
        print(f"Total received power shape: {total_mr.shape}")
        print(f"Post-combining noise shape: {noise_mr.shape}")

        print(
            "SINR terms finite: "
            f"{np.isfinite(desired_mr).all() and np.isfinite(total_mr).all() and np.isfinite(noise_mr).all()}"
        )
                # Kiem tra ham tinh SINR va SE
        sinr_test, se_test = compute_uatf_sinr_and_se(
            desired_mr,
            total_mr,
            noise_mr,
            P_UE_W,
            tau_c,
            cfg["tau_p"]
        )

        print("SINR test shape:", sinr_test.shape)
        print("SE test shape:", se_test.shape)
        print("SINR test finite:", np.isfinite(sinr_test).all())
        print("SE test finite:", np.isfinite(se_test).all())
        
        print(f"MR combiner shape: {V_MR.shape}")
        print(f"ZF combiner shape: {V_ZF.shape}")
        print(f"M-MMSE combiner shape: {V_MMMSE.shape}")

        print(
            "All combiners finite: "
            f"{np.isfinite(V_MR).all() and np.isfinite(V_ZF).all() and np.isfinite(V_MMMSE).all()}"
        )

        # Kiem tra ket qua uoc luong
        nmse = (
            np.mean(np.abs(H - H_hat) ** 2)
            / np.mean(np.abs(H) ** 2)
        )

        print(f"MMSE channel estimate shape: {H_hat.shape}")
        print(f"MMSE estimate finite: {np.isfinite(H_hat).all()}")
        print(f"Normalized MSE: {nmse:.4f}")

        print(f"Received pilot matrix shape: {Y_p.shape}")
        print(f"Pilot observation shape: {y_obs.shape}")
        print(
            f"Pilot observations finite: "
            f"{np.isfinite(y_obs).all()}"
        )
            # Chay mo phong Monte Carlo day du
    print("\nBAT DAU MO PHONG MONTE CARLO")

    mc_results = run_monte_carlo_simulation(
        noise_power_W
    )

    print("\nMO PHONG MONTE CARLO HOAN TAT")
    export_results_to_excel(mc_results)


# ==========================================================
# 7. CHAY CHUONG TRINH
# ==========================================================

if __name__ == "__main__":
    main()